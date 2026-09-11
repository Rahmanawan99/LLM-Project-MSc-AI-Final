import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document

doc = Document(r'c:\Users\rahma\Documents\Playground\LLM-Project\DMUD_MSC_Thesis-Final_v1.3.docx')

print("=== IMAGES IN DOCUMENT AND THEIR SURROUNDINGS ===")
for i, p in enumerate(doc.paragraphs):
    for r in p.runs:
        if 'blip' in r._element.xml:
            # find blip r:embed
            import xml.etree.ElementTree as ET
            tree = ET.fromstring(r._element.xml)
            # find any blip
            blips = [elem.attrib for elem in tree.iter() if elem.tag.endswith('blip')]
            surround = ""
            for k in range(max(0, i-2), min(len(doc.paragraphs), i+3)):
                t = doc.paragraphs[k].text.strip()
                if t:
                    surround += f"[{k}: {t[:40]}] "
            print(f"Para [{i}]: blips={blips} -> Context: {surround}")
