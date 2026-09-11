import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document

doc = Document(r'c:\Users\rahma\Documents\Playground\LLM-Project\DMUD_MSC_Thesis-Final_v1.3.docx')

print(f"Total tables: {len(doc.tables)}")
for idx, tbl in enumerate(doc.tables):
    headers = [c.text.strip().replace('\n', ' ') for c in tbl.rows[0].cells]
    # Check text around table element
    parent = tbl._element.getparent()
    pos = list(parent).index(tbl._element)
    prec_text = ""
    for k in range(max(0, pos-2), pos):
        el = parent[k]
        if hasattr(el, 'text') and el.text:
            prec_text += el.text.strip() + " | "
    print(f"Table #{idx+1} (Rows: {len(tbl.rows)}, Cols: {len(tbl.columns)})")
    print(f"  Preceding: {prec_text[:80]}")
    print(f"  Header: {headers[:4]}")
