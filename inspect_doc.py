import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document

doc = Document(r'c:\Users\rahma\Documents\Playground\LLM-Project\DMUD_MSC_Thesis-Final_v1.2.docx')

print('=== PARAGRAPHS 198 to 240 ===')
for i in range(198, min(240, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    print(f'[{i}] ({p.style.name if p.style else None}) "{p.text.strip()}"')

print('\n=== PARAGRAPHS 300 to 326 (Section 5.7) ===')
for i in range(300, min(326, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    print(f'[{i}] ({p.style.name if p.style else None}) "{p.text.strip()}"')
