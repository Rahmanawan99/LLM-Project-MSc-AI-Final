import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
import os

doc = Document(r'c:\Users\rahma\Documents\Playground\LLM-Project\DMUD_MSC_Thesis-Final_v1.3.docx')

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

# Check embedded images
images = [rel.target_ref for rel in doc.part.rels.values() if 'image' in rel.target_ref]
print(f"Total Embedded Images: {len(images)}")
for img in images:
    print(f"  - {img}")

# Check headings and structure
with open(r'c:\Users\rahma\Documents\Playground\LLM-Project\v1_3_dump.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        style = p.style.name if p.style else 'None'
        text = p.text.strip()
        f.write(f"[{i}] ({style}) {text}\n")
        if 'Heading' in style or 'Caption' in style:
            print(f"[{i}] ({style}) {text[:90]}")

print("\nDumped paragraphs to v1_3_dump.txt")
