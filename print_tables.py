import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document

doc = Document(r'c:\Users\rahma\Documents\Playground\LLM-Project\DMUD_MSC_Thesis-Final_v1.3.docx')

for i in range(3, len(doc.tables)):
    tbl = doc.tables[i]
    print(f"\n=================== TABLE #{i+1} ===================")
    for r in tbl.rows:
        print(" | ".join([c.text.strip().replace('\n', ' ') for c in r.cells]))
