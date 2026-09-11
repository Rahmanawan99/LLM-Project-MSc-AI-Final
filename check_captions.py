import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document

doc = Document(r'c:\Users\rahma\Documents\Playground\LLM-Project\DMUD_MSC_Thesis-Final_v1.3.docx')

print("=== TABLE 1: LIST OF FIGURES (Front Matter) ===")
for r in doc.tables[0].rows:
    print([c.text.strip().replace('\n', ' ') for c in r.cells])

print("\n=== TABLE 2: LIST OF TABLES (Front Matter) ===")
for r in doc.tables[1].rows:
    print([c.text.strip().replace('\n', ' ') for c in r.cells])

print("\n=== APPENDIX CAPTIONS & SURROUNDING TEXT ===")
for i in range(438, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    style = p.style.name if p.style else 'None'
    print(f"[{i}] ({style}) {p.text.strip()}")
